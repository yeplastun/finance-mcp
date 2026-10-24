"""A one-way copy of the ledger in Google Drive: one Google Sheet, rewritten after every ledger update.

The token is drive.file only (GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET and GOOGLE_REFRESH_TOKEN in state/env,
from an OAuth client of your own; see README.md), so it reaches only files this OAuth client
created. The workbook is built as .xlsx and Drive converts it into the Sheet, replacing what was there; the
Sheet's ID stays in state/sheet.json, so its link never changes. Edits made in the Sheet are overwritten:
labels change through the finance MCP. Text is always stored as text, never as a formula, because bank
descriptions are written by other people and "=IMPORTXML(…)" in one must not run.
"""

import io
import json
import secrets
from collections.abc import Sequence
from decimal import Decimal
from zoneinfo import ZoneInfo

import httpx
import psycopg
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

TOKEN_URL = "https://oauth2.googleapis.com/token"
UPLOAD = "https://www.googleapis.com/upload/drive/v3/files"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET = "application/vnd.google-apps.spreadsheet"
NAME = "Finance ledger"
LONDON = ZoneInfo("Europe/London")
# (heading, width, number format)
COLUMNS = (
    ("Ref", 8, "0"),
    ("Date", 11, "yyyy-mm-dd"),
    ("Amount £", 11, "#,##0.00"),
    ("Merchant", 28, None),
    ("Category", 22, None),
    ("Flow", 12, None),
    ("Person", 14, None),
    ("Whose", 10, None),
    ("Account", 22, None),
    ("Trip", 20, None),
    ("Tags", 16, None),
    ("Note", 24, None),
    ("Description", 40, None),
    ("Status", 9, None),
    ("Added", 16, "yyyy-mm-dd hh:mm"),
)
QUERY = """select num, day, pence, coalesce(merchant, description), category, flow, person, whose,
    initcap(replace(provider, 'ob-', '')) || ' ' || account, trip, array_to_string(tags, ', '), note,
    description, status, added
    from ledger.entry order by day desc, at desc, num desc"""


class SheetError(Exception):
    """Safe to display: never a response body or token."""


def workbook(rows: Sequence[tuple]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Entries"
    sheet.append([heading for heading, _, _ in COLUMNS])
    for num, day, pence, *text, status, added in rows:
        amount = Decimal(pence).scaleb(-2)
        sheet.append([num, day, amount, *text, status, added.astimezone(LONDON).replace(tzinfo=None)])
        for cell in sheet[sheet.max_row]:
            if cell.data_type == "f":  # openpyxl takes any text starting with "=" for a formula
                cell.data_type = "s"
    for index, (_, width, number_format) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        if number_format:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = number_format
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def access_token(http: httpx.Client, google: dict[str, str]) -> str:
    response = http.post(
        TOKEN_URL,
        data={
            "client_id": google["CLIENT_ID"],
            "client_secret": google["CLIENT_SECRET"],
            "refresh_token": google["REFRESH_TOKEN"],
            "grant_type": "refresh_token",
        },
    )
    if response.status_code != 200:
        raise SheetError(
            f"Google token refresh failed (HTTP {response.status_code}); rerun drive_token.py --for finance"
        )
    return response.json()["access_token"]


def upload(http: httpx.Client, token: str, data: bytes, sheet_id: str | None) -> str:
    """Replace the Sheet's contents, or create it the first time (or after it was deleted)."""
    auth = {"Authorization": f"Bearer {token}"}
    if sheet_id:
        response = http.patch(
            f"{UPLOAD}/{sheet_id}",
            params={"uploadType": "media", "fields": "id"},
            headers={**auth, "Content-Type": XLSX},
            content=data,
        )
        if response.status_code == 200:
            return sheet_id
        if response.status_code != 404:
            raise SheetError(f"Drive refused the update (HTTP {response.status_code})")
    boundary = secrets.token_hex(16)
    meta = json.dumps({"name": NAME, "mimeType": SHEET})
    body = (
        (
            f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{meta}\r\n"
            f"--{boundary}\r\nContent-Type: {XLSX}\r\n\r\n"
        ).encode()
        + data
        + f"\r\n--{boundary}--\r\n".encode()
    )
    response = http.post(
        UPLOAD,
        params={"uploadType": "multipart", "fields": "id"},
        headers={**auth, "Content-Type": f"multipart/related; boundary={boundary}"},
        content=body,
    )
    if response.status_code != 200:
        raise SheetError(f"Drive refused to create the Sheet (HTTP {response.status_code})")
    return response.json()["id"]


def export(
    conn: psycopg.Connection, http: httpx.Client, google: dict[str, str], sheet_id: str | None
) -> tuple[str, int]:
    """Rewrite the Sheet from ledger.entry. Returns its ID (new if it had to be created) and the entry count."""
    rows = conn.execute(QUERY).fetchall()
    return upload(http, access_token(http, google), workbook(rows), sheet_id), len(rows)
